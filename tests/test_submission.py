import hashlib
import zipfile
from pathlib import Path

import pandas as pd
from docx import Document


ROOT = Path(__file__).resolve().parents[1]


def test_verified_references_have_unique_keys_and_urls() -> None:
    references = pd.read_csv(ROOT / "references_verified.csv")
    assert references["key"].is_unique
    assert references["url"].str.startswith("http").all()
    assert references["verification_source"].notna().all()


def test_manuscript_contains_all_cited_figures_and_tables() -> None:
    document = Document(ROOT / "manuscript/manuscript.docx")
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert len(document.inline_shapes) == 5
    assert len(document.tables) == 3
    for number in range(1, 6):
        assert f"Fig. {number}" in text
    for number in range(1, 4):
        assert f"Table {number}" in text
    abstract = next(
        document.paragraphs[index + 1].text
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text == "Abstract"
    )
    assert 150 <= len(abstract.split()) <= 250
    keywords = next(
        paragraph.text.removeprefix("Keywords: ")
        for paragraph in document.paragraphs
        if paragraph.text.startswith("Keywords: ")
    )
    assert len(keywords.split("; ")) == 6


def test_submission_archive_and_manifest_exist() -> None:
    archive_path = ROOT / "submission/journal_of_geodesy_submission_FINAL.zip"
    assert archive_path.stat().st_size > 0
    manifest = pd.read_csv(ROOT / "submission_final/MANIFEST.csv")
    assert len(manifest) >= 20
    assert manifest["sha256"].str.len().eq(64).all()
    for row in manifest.itertuples():
        path = ROOT / "submission_final" / row.file
        assert path.stat().st_size == row.size_bytes
        assert hashlib.sha256(path.read_bytes()).hexdigest() == row.sha256
    with zipfile.ZipFile(archive_path) as archive:
        assert archive.testzip() is None
        assert set(archive.namelist()) == set(manifest["file"]) | {"MANIFEST.csv"}
        assert "manuscript_inline.docx" in archive.namelist()
        assert "DATA_CODE_AVAILABILITY.txt" in archive.namelist()
        assert "AI_DISCLOSURE.txt" in archive.namelist()
        assert (
            "figures/figure_2_residual_invariance_hierarchy.png"
            in archive.namelist()
        )
        assert "figures/figure_2_stabilizer_hierarchy.png" not in archive.namelist()
        assert "figures/figure_5_itrf_demonstration.png" in archive.namelist()
        assert "figures/supplementary_figure_s2_anchor_sensitivity.png" in archive.namelist()
        assert "figures/figure_6_itrf_demonstration.png" not in archive.namelist()
    with zipfile.ZipFile(ROOT / "submission/dtrs_reproducibility.zip") as archive:
        assert archive.testzip() is None
        names = set(archive.namelist())
        assert "data/raw/itrf2020/ITRF2020_SLR.SSC.txt" in names
        assert (
            "data/raw/literature/crossref_asimow_roth_1978_2026-09-25.json"
            in names
        )
        assert "scripts/run_pipeline.py" in names
        assert "scripts/build_submission.py" in names
        assert "scripts/verify_submission.py" in names
