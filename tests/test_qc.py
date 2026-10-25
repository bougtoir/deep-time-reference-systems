"""Final-QC guards: OMML math, citation completeness/order, markup, grid,
hard-coded results, and traceability currency."""
import re
import zipfile
from pathlib import Path

import pandas as pd
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCX_FILES = [
    ROOT / "manuscript/manuscript.docx",
    ROOT / "supplement/supplement.docx",
    ROOT / "manuscript/tables.docx",
    ROOT / "submission/cover_letter.docx",
    ROOT / "submission/title_page.docx",
    ROOT / "submission/declarations.docx",
]


def _document_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        return archive.read("word/document.xml").decode("utf-8")


def _body_text(document: Document) -> str:
    return "\n".join(p.text for p in document.paragraphs)


# A. OMML math is present and covers the key equations
def test_manuscript_math_is_omml() -> None:
    xml = _document_xml(ROOT / "manuscript/manuscript.docx")
    assert xml.count("<m:oMath>") > 0
    math_text = "".join(re.findall(r"<m:t[^>]*>([^<]*)</m:t>", xml))
    for fragment in ["arcsin", "atan2", "det", "π", "ℝ", "∘"]:
        assert fragment in math_text, fragment


# B. Every supplementary item is cited in manuscript prose
def test_supplementary_items_are_cited() -> None:
    text = _body_text(Document(ROOT / "manuscript/manuscript.docx"))
    for label in [
        "Supplementary Fig. S1",
        "Supplementary Fig. S2",
        "Supplementary Table S1",
        "Supplementary Table S2",
        "Supplementary Table S3",
    ]:
        assert label in text, label


# C. First prose citations are monotonic within each numbered family
def test_first_citation_order_is_monotonic() -> None:
    text = _body_text(Document(ROOT / "manuscript/manuscript.docx"))
    families = [
        ["Fig. 1", "Fig. 2", "Fig. 3", "Fig. 4", "Fig. 5"],
        ["Table 1", "Table 2", "Table 3"],
        ["Supplementary Table S1", "Supplementary Table S2",
         "Supplementary Table S3"],
        ["Supplementary Fig. S1", "Supplementary Fig. S2"],
    ]
    for family in families:
        positions = [text.find(label) for label in family]
        assert all(p >= 0 for p in positions), (family, positions)
        assert positions == sorted(positions), (family, positions)


# D. Every bibliography entry is cited; every author-year citation resolves
def test_references_cited_and_resolved() -> None:
    document = Document(ROOT / "manuscript/manuscript.docx")
    paragraphs = [p.text for p in document.paragraphs]
    ref_start = paragraphs.index("References")
    body = "\n".join(paragraphs[:ref_start])
    references = [p for p in paragraphs[ref_start + 1:] if p.strip()]
    assert len(references) == 24
    for reference in references:
        surname = reference.split(",")[0].split(" ")[0]
        year = re.search(r"\((\d{4})", reference) or re.search(r"\b(19|20)\d{2}\b", reference)
        year = year.group(1) if year else None
        if surname == "ISO" or reference.startswith("International Organization"):
            assert "ISO 2019" in body
            continue
        assert surname in body and year in body, reference
    for group in re.findall(r"\(([^()]*\b(?:19|20)\d{2}[a-z]?[^()]*)\)", body):
        if "http" in group:
            continue
        for part in group.split(";"):
            part = part.strip()
            if not re.search(r"(?:19|20)\d{2}", part):
                continue
            if not re.match(r"^[A-Z][A-Za-zÀ-ž'\-]+(\s+(and|et al\.?)\s+|\s)\S*", part):
                continue
            surname = re.split(r"[ ,]", part)[0]
            year = re.search(r"((?:19|20)\d{2})", part).group(1)
            resolved = any(
                surname in reference and year in reference
                for reference in references
            )
            assert resolved, part


# E. Headline results are not hard-coded outside generated registries
def test_headline_results_not_hardcoded() -> None:
    headline = [
        "0.9465", "0.947", "0.996", "0.691", "0.091", "0.000225",
        "0.000121", "6136173", "6.136", "0.895", "0.602",
        "0.9358", "0.9555", "0.9921", "0.9980",
    ]
    allowed = {
        ROOT / "results/manuscript_values.csv",
        ROOT / "result_traceability.csv",
        ROOT / "results/rigidity_phase_diagram.csv",
        ROOT / "results/ellipsoid_sensitivity.csv",
        ROOT / "results/distance_reconstruction.csv",
        ROOT / "results/itrf_distance_changes.csv",
    }
    scanned = list((ROOT / "src").rglob("*.py")) + list(
        (ROOT / "scripts").glob("*.py")
    ) + list((ROOT / "configs").glob("*.yml"))
    violations = []
    for path in scanned:
        text = path.read_text(encoding="utf-8")
        for value in headline:
            if value in text:
                violations.append(f"{path.relative_to(ROOT)}: {value}")
    assert not violations, violations


# F. No Markdown asterisk artifacts in any submission DOCX
def test_no_markdown_asterisks() -> None:
    for path in DOCX_FILES:
        document = Document(path)
        text = _body_text(document)
        assert "**" not in text, path
        for line in text.split("\n"):
            assert not line.strip().startswith(("# ", "* ")), (path, line[:40])
        # single '*' used as markup would wrap a word
        assert not re.search(r"\*[A-Za-z][^*]*\*", text), path


# G. No character-count document grid or forced wrapping
def test_no_character_grid() -> None:
    for path in DOCX_FILES:
        xml = _document_xml(path)
        assert "w:docGrid" not in xml, path
        assert 'w:type="linesAndChars"' not in xml, path
        assert "w:charSpace" not in xml, path


# H. Traceability labels match the canonical supplement numbering
def test_traceability_numbering_is_current() -> None:
    frame = pd.read_csv(ROOT / "result_traceability.csv")
    canonical = {
        "Supplementary Table S1",
        "Supplementary Table S2",
        "Supplementary Table S3",
    }
    table_labels = set()
    for value in frame["figure_or_table"]:
        for label in str(value).split(";"):
            label = label.strip()
            if label.startswith("Supplementary Table"):
                table_labels.add(label)
    assert table_labels <= canonical
    ellipsoid = frame[frame["claim_id"] == "C6"].iloc[0]
    assert ellipsoid["figure_or_table"] == "Supplementary Table S2"


def test_manuscript_values_registry_exists_and_nonempty() -> None:
    frame = pd.read_csv(ROOT / "results/manuscript_values.csv")
    assert len(frame) >= 20
    assert frame["metric"].is_unique
