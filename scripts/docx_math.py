"""Insert native Word equations (OMML) into python-docx paragraphs.

Prose strings embed LaTeX inside ``$...$`` markers. ``add_rich_paragraph``
splits the string, emits ordinary runs for prose, and inserts an OMML
``<m:oMath>`` element for each math span so equations stay editable in
Microsoft Word.
"""
import re

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from latex2mathml.converter import convert as latex_to_mathml
from docx_equation import mathml_to_omml

MATH_RE = re.compile(r"\$([^$\n]+)\$")


def insert_omml(paragraph, latex: str) -> None:
    """Append an OMML rendering of ``latex`` to ``paragraph``."""
    omml = mathml_to_omml(latex_to_mathml(latex))
    for index, run in enumerate(omml.iter(qn("m:r"))):
        text = run.find(qn("m:t"))
        if text is not None and (
            text.text == "*"
            or (index == 0 and text.text in {"<", ">", "≤", "≥"})
        ):
            rpr = run.find(qn("m:rPr"))
            if rpr is None:
                rpr = OxmlElement("m:rPr")
                run.insert(0, rpr)
            literal = OxmlElement("m:lit")
            literal.set(qn("m:val"), "1")
            rpr.insert(0, literal)
    paragraph._element.append(omml)


def add_rich_runs(paragraph, text: str) -> None:
    """Append ``text`` to ``paragraph``, converting ``$...$`` spans to OMML."""
    position = 0
    for match in MATH_RE.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        insert_omml(paragraph, match.group(1))
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def add_rich_paragraph(document, text: str):
    """``document.add_paragraph`` equivalent that converts ``$...$`` to OMML."""
    paragraph = document.add_paragraph()
    add_rich_runs(paragraph, text)
    return paragraph
